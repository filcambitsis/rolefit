import io
import re
import zipfile

from docx import Document

from .normalization import minimum_years
from .skills import mentions, normalize_skill

# Bump this when extraction rules or the skills vocabulary change:
# stored requirements with an older version are extracted again.
PROMPT_VERSION = "rules-v11"
MAX_TEXT = 150_000


def parse_file(data: bytes, filename: str) -> str:
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("CV must be 5 MB or smaller")
    if filename.lower().endswith(".pdf"):
        import pymupdf

        if not data.startswith(b"%PDF"):
            raise ValueError("This file is not a valid PDF")
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.page_count > 30:
                raise ValueError("CV must be 30 pages or fewer")
            texts = []
            for page in doc:
                blocks = [b for b in page.get_text("blocks") if b[6] == 0]
                mid = page.rect.width / 2
                left = [b for b in blocks if b[2] <= mid + 10]
                right = [b for b in blocks if b[0] >= mid - 10]
                if len(left) >= 2 and len(right) >= 2:
                    header = [b for b in blocks if b not in left and b not in right]
                    ordered = (
                        sorted(header, key=lambda b: b[1])
                        + sorted(left, key=lambda b: b[1])
                        + sorted(right, key=lambda b: b[1])
                    )
                else:
                    ordered = sorted(blocks, key=lambda b: (round(b[1] / 8), b[0]))
                texts.append("\n".join(b[4].strip() for b in ordered))
            text = "\n\n".join(texts)
    elif filename.lower().endswith(".docx"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(z.file_size for z in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError("Expanded document is too large")
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    elif filename.lower().endswith(".txt"):
        text = data.decode("utf-8-sig")
    else:
        raise ValueError("Supported formats: PDF, DOCX, TXT")
    if len(text.strip()) < 50:
        raise ValueError("Not enough readable text. Scanned PDFs need OCR before upload.")
    if len(text) > MAX_TEXT:
        raise ValueError("CV text is too long")
    return text


def verified_span(raw: str, quote: str):
    start = raw.find(quote)
    if not quote.strip() or start < 0:
        return None
    return {"quote": quote, "start": start, "end": start + len(quote)}


SECTION_HEADINGS = {
    "experience": "experience",
    "work experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "education": "education",
    "education and qualifications": "education",
    "projects": "projects",
    "key projects": "projects",
    "personal projects": "projects",
    "skills": "skills",
    "technical skills": "skills",
    "skills and interests": "skills",
    "languages": "languages",
    "profile": "other",
    "summary": "other",
    "additional information": "other",
    "leadership and activities": "other",
}


def extract_cv(raw):
    candidates, section = [], "other"
    for line in raw.splitlines():
        quote = line.strip()
        heading = " ".join(quote.lower().replace("&", "and").rstrip(":").split())
        if heading in SECTION_HEADINGS:
            section = SECTION_HEADINGS[heading]
        elif len(quote) >= 15 or mentions(quote):
            candidates.append({"quote": quote, "section": section})
    evidence, failures = [], 0
    for item in candidates:
        span = verified_span(raw, item["quote"])
        if span is None:
            failures += 1
        else:
            evidence.append({**span, "section": item["section"], "skills": mentions(item["quote"])})
    return evidence, failures, len(candidates)


def extract_requirements(raw):
    candidates, preferred, in_requirements = [], False, False
    for line in raw.splitlines():
        line = line.strip()
        heading = re.sub(r"[-–]", " ", line.lower()).strip(": ")
        # Section boundaries prevent company blurbs and benefits becoming requirements.
        if len(line) < 160 and re.search(
            r"^(nice to haves?|preferred qualifications|preferred requirements|bonus|desirable|"
            r"especially strong backgrounds|it would be great|while it.s not required|any of the following|"
            r"ideally you.{0,3}(?:d |would )have)",
            heading,
        ):
            preferred, in_requirements = True, True
            continue
        if len(line) < 100 and re.search(
            r"^(minimum requirements|required qualifications|requirements|qualifications|must haves?|what you.bring|"
            r"it.s important to us|essential skills|minimum qualifications|"
            r"what you.ll bring|what we.re looking for|you may be a fit if|you should have|you.ll need|"
            r"we.d love|we.re looking for|you might be a fit|skills and experience|what you.ll need|"
            r"what you will (?:need|bring)|what you need|who you are|about you|what we look for|skills you|you might thrive|you may be a good fit)",
            heading,
        ):
            preferred, in_requirements = False, True
            continue
        if re.search(
            r"^(about us|about the |who we are|what you.ll do|you will:?$|responsibilities|"
            r"what we offer|what we expect|recruitment steps|meet your team|after you apply|benefits|compensation|applying|please note|we offer|equal opportunity|about |"
            r"we hire|full.time employees|how and where we work|a note on ai|by clicking|#li[ -]|notice$|working location|our research interviews)",
            heading,
        ):
            in_requirements = False
            continue
        if re.search(
            r"equal opportunity|privacy policy|compensation offered|cash compensation|"
            r"reasonable accommodations|base salary|we encourage you to apply",
            line,
            re.I,
        ):
            in_requirements = False
            continue
        if len(line) < 5 or len(line) > 1500:
            continue
        if not in_requirements and not re.search(
            r"^(?:you must|we require|must have)",
            line,
            re.I,
        ):
            continue
        if re.search(r"no .{0,30}(?:experience|degree) (?:is )?(?:required|necessary)", line, re.I):
            continue
        required = not preferred and not bool(
            re.search(r"nice to have|(?:is|will be) a plus|not required|preferred", line, re.I)
        )
        # Preserve unsupported requirements too; dropping them inflates coverage.
        if re.search(r"\b(?:bachelor|master|ph\.?d|ph\.d\.|BS/BA|MS/MA|BSc|MSc|BS|MS|degree)\b", line, re.I):
            candidates.append({"quote": line, "category": "education", "skill": None, "required": required})
        elif minimum_years(line) is not None:
            candidates.append({"quote": line, "category": "experience", "skill": None, "required": required})
        else:
            skills = mentions(line)
            if skills and re.search(r"\bor\b", line, re.I):
                candidates.append({"quote": line, "category": "other", "skill": None, "required": required})
            elif skills:
                candidates.extend(
                    {"quote": line, "category": "skill", "skill": skill, "required": required}
                    for skill in skills
                )
            elif in_requirements:
                candidates.append({"quote": line, "category": "other", "skill": None, "required": required})
    out, seen = [], set()
    for candidate in candidates:
        if verified_span(raw, candidate["quote"]) is None:
            continue
        skill = normalize_skill(candidate["skill"])[0] if candidate["skill"] else None
        # A normalized skill cannot be introduced solely by model suggestion.
        if skill and skill not in mentions(candidate["quote"]):
            continue
        identity = (skill, candidate["quote"], candidate["required"])
        if identity in seen:
            continue
        seen.add(identity)
        out.append(
            {
                "quote": candidate["quote"],
                "text": candidate["quote"],
                "category": candidate["category"],
                "skill": skill,
                "required": candidate["required"],
                "min_years": minimum_years(candidate["quote"]),
            }
        )
    return out
